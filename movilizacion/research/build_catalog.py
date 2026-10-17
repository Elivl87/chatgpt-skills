import csv, sys, glob, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

R = sys.argv[1]; OUT = sys.argv[2]
order = ["excel","ppt","finanzas","riesgo","oilgas"]
rows = {}
prank = {"Alta":0,"Media":1,"Baja":2}
for f in order:
    with open(os.path.join(R, f + ".tsv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh, delimiter="\t"):
            k = r["repo"].lower()
            if k in rows:
                e = rows[k]
                if r["categoria"] not in e["cats"]: e["cats"].append(r["categoria"])
                if prank[r["prioridad"]] < prank[e["prioridad"]]: e["prioridad"] = r["prioridad"]
                if r["como_ayuda"] not in e["usos"]: e["usos"].append(r["como_ayuda"])
            else:
                r["cats"] = [r["categoria"]]; r["usos"] = [r["como_ayuda"]]
                rows[k] = r

# Proposed stack: repo key -> (layer, role)
stack = {
 "jmcnamara/xlsxwriter": ("1. Construir el Excel", "Escribir el libro con fórmulas vivas, nombres, validaciones y gráficos nativos"),
 "openpyxl": ("1. Construir el Excel", "Releer, verificar y editar el libro (ya instalado)"),
 "anthropics/skills (xlsx/recalc.py)": ("2. Validar el Excel", "Recalcular con LibreOffice y garantizar cero errores (ya disponible)"),
 "vinci1it2000/formulas": ("2. Validar el Excel", "Segundo motor: recalcular sin Excel y probar celdas clave"),
 "anthropics/financial-services": ("2. Validar el Excel", "Checklist de auditoría (audit-xls) y convenciones de modelado"),
 "numpy/numpy-financial": ("3. Evaluación económica", "Verificar VPN/TIR si se usan flujos en el tiempo"),
 "phIlipgriffith/ahpy".lower(): ("3. Evaluación económica", "Pesos de criterios cualitativos con consistencia (se replica en Excel)"),
 "nlrwindsystems/landbosse": ("3. Evaluación económica", "Referencia de rubros: transporte, grúas, gestión, demoras por clima"),
 "juansecorreapaez/riggo": ("3. Evaluación económica", "Estructura del inventario de cargas (solo estructura; sin licencia)"),
 "praveen-christuraj/well-costing": ("3. Evaluación económica", "Estructura de rate card hora/día/único (solo estructura; sin licencia)"),
 "scottshambaugh/monaco": ("4. Riesgo y sensibilidad", "Monte Carlo de costo y plazo: P10/P50/P90 y P(B < A)"),
 "tadamcz/betapert": ("4. Riesgo y sensibilidad", "Distribuciones PERT de 3 puntos (mín/probable/máx)"),
 "salib/salib": ("4. Riesgo y sensibilidad", "Ranking de variables que mueven la diferencia A − B"),
 "netflix-skunkworks/riskquant": ("4. Riesgo y sensibilidad", "Patrón para convertir el registro de riesgos en dinero"),
 "pyscripter/xlrisk": ("4. Riesgo y sensibilidad", "Opcional: que tú corras el Monte Carlo dentro de tu Excel de escritorio"),
 "scanny/python-pptx": ("5. Presentación", "Generar el deck con gráficos nativos editables"),
 "anthropics/skills": ("5. Presentación", "Skill pptx de Claude: creación y QA visual (ya disponible)"),
 "likaku/mck-ppt-design-skill": ("5. Presentación", "Patrones de diseño estilo consultora para gerencia"),
 "matplotlib/matplotlib": ("5. Presentación", "Matriz de riesgos 5×5 e histogramas como imagen"),
 "libreoffice/core": ("6. Control de calidad", "Convertir pptx/xlsx a PDF para revisión"),
 "belval/pdf2image": ("6. Control de calidad", "PDF a imágenes para revisar cada diapositiva"),
 "microsoft/markitdown": ("6. Control de calidad", "Cruzar cifras del deck contra el Excel"),
}

data = sorted(rows.values(), key=lambda r: (prank[r["prioridad"]], r["cats"][0], r["repo"].lower()))

F = "Arial"
hdr_font = Font(name=F, bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="1F3864")
base = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
title = Font(name=F, size=14, bold=True, color="1F3864")
yellow = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
# --- Catálogo
ws = wb.active; ws.title = "Catalogo"
cols = [("ID",6),("Repositorio",34),("URL",40),("Categoría(s)",22),("Lenguaje",16),("Estrellas",10),
        ("Última actividad",18),("Licencia",16),("Qué hace",45),("Cómo ayuda a la evaluación",55),
        ("Prioridad",10),("En pila propuesta",12),("Capa (propuesta)",24),("Rol en la pila",45),
        ("Nota",30),("Tu decisión",12),("Comentarios",35)]
ws["A1"] = "Catálogo de repositorios GitHub – Evaluación de movilización (lump sum vs. tercerizado)"; ws["A1"].font = title
ws["A2"] = "Edita solo las columnas amarillas (Tu decisión: Sí / No / Revisar; Comentarios). Fuente: 5 agentes buscadores, 12-oct-2026; estrellas y fechas verificadas en GitHub/PyPI."
ws["A2"].font = Font(name=F, size=9, italic=True)
HR = 4
for i,(h,w) in enumerate(cols,1):
    c = ws.cell(HR, i, h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = box
    ws.column_dimensions[get_column_letter(i)].width = w
for n, r in enumerate(data, 1):
    rr = HR + n
    key = r["repo"].lower()
    st = stack.get(key)
    vals = [n, r["repo"], r["url"], "; ".join(r["cats"]), r["lenguaje"], r["estrellas"], r["ultima_actividad"],
            r["licencia"], r["que_hace"], " | ".join(r["usos"]), r["prioridad"], "Sí" if st else "No",
            st[0] if st else "", st[1] if st else "", r.get("nota") or "", "", ""]
    for i, v in enumerate(vals, 1):
        c = ws.cell(rr, i, v); c.font = base; c.alignment = wrap; c.border = box
    ws.cell(rr, 3).hyperlink = r["url"]; ws.cell(rr, 3).font = Font(name=F, size=10, color="0563C1", underline="single")
    ws.cell(rr, 16).fill = yellow; ws.cell(rr, 17).fill = yellow
last = HR + len(data)
dv = DataValidation(type="list", formula1='"Sí,No,Revisar"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"P{HR+1}:P{last}")
ws.auto_filter.ref = f"A{HR}:Q{last}"; ws.freeze_panes = f"C{HR+1}"
rng = f"K{HR+1}:K{last}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Alta"'], fill=PatternFill("solid", fgColor="C6EFCE")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Media"'], fill=PatternFill("solid", fgColor="FFEB9C")))
ws.conditional_formatting.add(f"L{HR+1}:L{last}", CellIsRule(operator="equal", formula=['"Sí"'], fill=PatternFill("solid", fgColor="BDD7EE"), font=Font(name=F, bold=True)))

# --- Pila propuesta
ps = wb.create_sheet("Pila propuesta")
ps["A1"] = "Pila de herramientas propuesta (para tu aprobación)"; ps["A1"].font = title
ps["A2"] = "Los datos se leen del Catálogo con fórmulas. Marca tu decisión en la hoja Catalogo, columna P."; ps["A2"].font = Font(name=F, size=9, italic=True)
heads = [("Capa",26),("Repositorio",34),("Rol en la pila",55),("Licencia",16),("Prioridad",10),("Tu decisión",12)]
for i,(h,w) in enumerate(heads,1):
    c = ps.cell(4, i, h); c.font = hdr_font; c.fill = hdr_fill; c.border = box
    ps.column_dimensions[get_column_letter(i)].width = w
items = sorted([(v[0], k) for k, v in stack.items()])
name_by_key = {r["repo"].lower(): r["repo"] for r in data}
for j, (layer, key) in enumerate(items, 5):
    repo = name_by_key[key]
    ps.cell(j, 1, layer); ps.cell(j, 2, repo)
    m = f'MATCH($B{j},Catalogo!$B${HR+1}:$B${last},0)'
    ps.cell(j, 3, f'=INDEX(Catalogo!$N${HR+1}:$N${last},{m})')
    ps.cell(j, 4, f'=INDEX(Catalogo!$H${HR+1}:$H${last},{m})')
    ps.cell(j, 5, f'=INDEX(Catalogo!$K${HR+1}:$K${last},{m})')
    ps.cell(j, 6, f'=IF(INDEX(Catalogo!$P${HR+1}:$P${last},{m})="","Pendiente",INDEX(Catalogo!$P${HR+1}:$P${last},{m}))')
    for i in range(1, 7):
        c = ps.cell(j, i); c.font = Font(name=F, size=10, color="008000") if i >= 3 else base; c.alignment = wrap; c.border = box
ps.freeze_panes = "A5"

# --- Resumen
rs = wb.create_sheet("Resumen", 0)
rs["A1"] = "Resumen de la búsqueda de repositorios"; rs["A1"].font = title
rs.column_dimensions["A"].width = 28
for col in "BCDEF": rs.column_dimensions[col].width = 14
rs["A3"] = "Categoría"; hs = ["Total","Alta","Media","Baja","En pila"]
for i, h in enumerate(["Categoría"] + hs, 1):
    c = rs.cell(3, i, h); c.font = hdr_font; c.fill = hdr_fill; c.border = box
cats = ["Excel","PowerPoint","Evaluación económica","Riesgo","Oil & Gas y logística"]
D = f"Catalogo!$D${HR+1}:$D${last}"; K = f"Catalogo!$K${HR+1}:$K${last}"; L = f"Catalogo!$L${HR+1}:$L${last}"
for j, cat in enumerate(cats, 4):
    rs.cell(j, 1, cat)
    crit = f'"*{cat}*"'
    rs.cell(j, 2, f'=COUNTIFS({D},{crit})')
    for i, p in enumerate(["Alta","Media","Baja"], 3):
        rs.cell(j, i, f'=COUNTIFS({D},{crit},{K},"{p}")')
    rs.cell(j, 6, f'=COUNTIFS({D},{crit},{L},"Sí")')
tr = 4 + len(cats)
rs.cell(tr, 1, "Repositorios únicos")
rs.cell(tr, 2, f'=COUNTA(Catalogo!$B${HR+1}:$B${last})')
for i, p in enumerate(["Alta","Media","Baja"], 3):
    rs.cell(tr, i, f'=COUNTIF({K},"{p}")')
rs.cell(tr, 6, f'=COUNTIF({L},"Sí")')
for j in range(4, tr + 1):
    for i in range(1, 7):
        c = rs.cell(j, i); c.font = bold if j == tr else base; c.border = box
rs.cell(tr + 1, 1, "Nota: un repo puede aparecer en varias categorías; por eso la suma por categoría supera el total único.").font = Font(name=F, size=9, italic=True)
dec = tr + 3
rs.cell(dec, 1, "Tus decisiones").font = bold
for k, lab in enumerate(["Sí","No","Revisar"]):
    rs.cell(dec + 1 + k, 1, lab).font = base
    rs.cell(dec + 1 + k, 2, f'=COUNTIF(Catalogo!$P${HR+1}:$P${last},"{lab}")').font = base
rs.cell(dec + 4, 1, "Pendientes").font = base
rs.cell(dec + 4, 2, f'=COUNTA(Catalogo!$B${HR+1}:$B${last})-SUM(B{dec+1}:B{dec+3})').font = base

h = dec + 6
rs.cell(h, 1, "Hallazgos clave").font = bold
notes = [
 "No existe en GitHub un modelo listo de 'movilización de taladro: lump sum vs. contratación directa'; se construye a medida con las piezas de la pila.",
 "No hay datasets públicos de tarifas de rig move ni de transporte pesado: tarifas y duraciones deben venir de tus datos (facturas, DDR, cotizaciones).",
 "Varios repos de dominio (RigGo, Well-Costing, Chaimurat) no tienen licencia: se usa solo su estructura como referencia, no su código.",
 "Las skills xlsx y pptx de Claude ya están disponibles en esta sesión; LibreOffice permite recalcular y revisar visualmente.",
 "XLRisk permitiría que tú mismo corras Monte Carlo dentro de tu Excel de escritorio (Windows/Mac).",
]
for k, t in enumerate(notes, 1):
    c = rs.cell(h + k, 1, f"• {t}"); c.font = base
    rs.merge_cells(start_row=h + k, start_column=1, end_row=h + k, end_column=6)
    c.alignment = Alignment(wrap_text=True, vertical="top"); rs.row_dimensions[h + k].height = 30
wb.save(OUT)
print(len(data), "repos únicos")
