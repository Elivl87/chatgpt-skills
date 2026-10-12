"""Lee los logs de movilizaciones (detalle por movimiento) y devuelve registros normalizados."""
from openpyxl import load_workbook

RANGOS = ["Misma locación", "0.15 - 5 km", "5 - 10 km", "10+ km"]


def leer(path):
    ws = load_workbook(path, data_only=True).worksheets[0]
    hdr = None
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row[0] == "Periodo":
            hdr = i
            continue
        if hdr and isinstance(row[0], int) and row[3] is not None:
            out.append({"fila": i, "anio": row[0], "n": row[1], "pozo": str(row[2]), "fecha": row[3],
                        "mes": row[4], "km": row[5], "horas": row[6], "rango": RANGOS.index(row[7])})
    return out
