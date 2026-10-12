"""Construye Auditoria_Evaluacion_Previa.xlsx desde el informe markdown de la auditoría (con cifras verificadas)."""
import sys, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule

MD, OUT = sys.argv[1], sys.argv[2]
txt = open(MD, encoding="utf-8").read()
# Correcciones verificadas contra el simulador (mismo motor): mezclas de los logs
FIX = {
    "Con el perfil de 2024, 15 + 20 en 0,1–5 km da PS 763.124 · EC 750.092 de ahorro. Con 15 + 20 repartidos según 2024 (12/6/2): PS 628.949 · EC 558.054. Con 15 + 20 según 2025 (10/6/4): PS 601.594 · EC 509.604.":
    "Simulador, perfil 2024: 15 + 20 en 0,1–5 km da PS 730.502 · EC 695.386. Con los 20 repartidos como 2024 (12/6/2): PS 615.695 · EC 532.934. Como 2025 (10/6/4): PS 599.214 · EC 502.719. Con perfil 2025 y mezcla 2025: PS 215.523 · EC 227.836.",
    "PS −134 k a −162 k; EC −192 k a −240 k (perfil 2024, mezcla 2024 o 2025)": "PS −115 k a −137 k; EC −162 k a −200 k (perfil 2024 o 2025, mezcla 2024 o 2025)",
    "Con la mezcla completa de 2024: PS −161 k": "No se cuantifica aparte: los 15 pozos son dato de SierraCol",
    "≈1.063 USD/h de equipos (flota 0–5 km de p.1 a tarifa horaria de p.5) más ≈161 USD/h de nómina (430 M COP / 720 h), en total **≈2.681 USD/h**":
    "≈1.357 USD/h de flota mínima (p.3) y 20 señalizadores a tarifa horaria de p.5, en total **≈2.813 USD/h** con el taladro (cifra única del simulador, hoja Pendientes)",
    "PS ≈ 6,8 h por movilización (18.095/2.681)": "PS ≈ 6,4 h por movilización (18.095/2.813)",
    "PS baja a ≈ 4,2 h": "PS baja a ≈ 4,0 h",
    "≈2.681 USD/h (PS) y ≥1.456 USD/h (EC). Unas 6,8 h por movilización anulan el ahorro de PS.": "≈2.813 USD/h (PS) y ≥1.456 USD/h (EC). Unas 6,4 h por movilización anulan el ahorro de PS.",
}
for a, b in FIX.items():
    assert a in txt, a[:60]
    txt = txt.replace(a, b)

def tables(md):
    out = []; cur = None; title = None
    for line in md.splitlines():
        if line.startswith("## "):
            title = line[3:].strip()
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if set("".join(cells)) <= set("-: "):
                continue
            if cur is None:
                cur = {"title": title, "header": cells, "rows": []}; out.append(cur)
            else:
                cur["rows"].append(cells)
        else:
            cur = None
    return out

def clean(s):
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s); s = s.replace("`", "")
    return s

T = tables(txt)
F = "Arial"; B = Font(name=F, size=10, bold=True); N = Font(name=F, size=10); NOTE = Font(name=F, size=9, italic=True, color="595959")
TITLE = Font(name=F, size=14, bold=True, color="1F3864"); H1 = Font(name=F, size=11, bold=True, color="FFFFFF")
F_H1 = PatternFill("solid", fgColor="1F3864"); F_H2 = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin); WR = Alignment(wrap_text=True, vertical="top")
COL = {"CONSERVAR": "C6EFCE", "CORREGIR": "FFEB9C", "DESCARTAR": "F8CBAD"}

wb = Workbook()
# ---- Veredictos
v = wb.active; v.title = "Veredictos"; v.sheet_view.showGridLines = False
for c, w in zip("ABCDEF", [2, 34, 26, 14, 14, 70]): v.column_dimensions[c].width = w
v["B1"] = "Auditoría de la evaluación previa (Zona 2 de Tarifas.xlsx)"; v["B1"].font = TITLE
v["B2"] = ("Se auditó después de construir el modelo propio. Cada fórmula de la Zona 2 se recalculó: la aritmética cuadra al centavo; los problemas están en supuestos, bases y textos. "
           "Veredicto: CONSERVAR / CORREGIR / DESCARTAR. Detalle con evidencia por celda en la hoja 'Detalle'.")
v["B2"].font = NOTE; v.merge_cells("B2:F2"); v["B2"].alignment = WR; v.row_dimensions[2].height = 30
hdr = ["Sección", "Celdas", "¿Bien hecha?", "¿Sirve?", "Veredicto y por qué"]
for i, h in enumerate(hdr):
    c = v.cell(4, 2 + i, h); c.font = B; c.fill = F_H2; c.border = BOX
r = 5
for t in T[:-1]:
    for row in t["rows"]:
        d = dict(zip(t["header"], row))
        ver = clean(d.get("Veredicto", ""))
        vals = [clean(d.get("Sección", "")) or t["title"], clean(d.get("Celdas", "")), clean(d.get("¿Bien hecha?", "")), clean(d.get("¿Sirve?", "")),
                ver + " — " + clean(d.get("Evidencia y corrección", ""))[:260] + ("…" if len(clean(d.get("Evidencia y corrección", ""))) > 260 else "")]
        for i, x in enumerate(vals):
            c = v.cell(r, 2 + i, x); c.font = N; c.border = BOX; c.alignment = WR
        key = next((k for k in COL if k in ver.upper()), None)
        if key: v.cell(r, 6).fill = PatternFill("solid", fgColor=COL[key])
        v.row_dimensions[r].height = 75
        r += 1
v.freeze_panes = "B5"
# ---- Detalle
d = wb.create_sheet("Detalle"); d.sheet_view.showGridLines = False
widths = [2, 26, 22, 18, 18, 14, 70, 22, 34]
for i, w in enumerate(widths): d.column_dimensions[chr(65 + i)].width = w
d["B1"] = "Detalle sección por sección: evidencia, su número y mi número"; d["B1"].font = TITLE
d["B2"] = "'Su número' = valor de la celda de la Zona 2. 'Mi número' = recálculo independiente (✓ = reproducido al centavo). Cifras de campaña verificadas contra el simulador."; d["B2"].font = NOTE
r = 4
for t in T[:-1]:
    d.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
    c = d.cell(r, 2, t["title"]); c.font = H1; c.fill = F_H1; r += 1
    for i, h in enumerate(t["header"]):
        c = d.cell(r, 2 + i, h); c.font = B; c.fill = F_H2; c.border = BOX; c.alignment = WR
    r += 1
    for row in t["rows"]:
        for i, x in enumerate(row):
            c = d.cell(r, 2 + i, clean(x)); c.font = N; c.border = BOX; c.alignment = WR
        d.row_dimensions[r].height = min(240, 15 * (1 + max(len(clean(x)) for x in row) // 70))
        r += 1
    r += 1
# ---- Sesgo
s = wb.create_sheet("Sesgo"); s.sheet_view.showGridLines = False
for c, w in zip("ABCDEF", [2, 5, 62, 16, 30, 46]): s.column_dimensions[c].width = w
s["B1"] = "¿Empujan todos los errores hacia el mismo lado?"; s["B1"].font = TITLE
t = T[-1]
for i, h in enumerate(t["header"]):
    c = s.cell(3, 2 + i, h); c.font = B; c.fill = F_H2; c.border = BOX; c.alignment = WR
for j, row in enumerate(t["rows"]):
    for i, x in enumerate(row):
        c = s.cell(4 + j, 2 + i, clean(x)); c.font = N; c.border = BOX; c.alignment = WR
    dirc = clean(row[3])
    s.cell(4 + j, 5).fill = PatternFill("solid", fgColor="FFEB9C" if "tercerizado" in dirc.lower() and "favorece tercerizado" in dirc.lower() else ("DDEBF7" if "WO" in dirc else "F2F2F2"))
    s.row_dimensions[4 + j].height = 42
rr = 5 + len(t["rows"])
concl = ("Conclusión: sí hay un sesgo sistemático. Los hallazgos de mayor peso (mezcla de rangos, horas por debajo de lo cotizado, frecuencia 5–7 tomada como 'rango real', "
         "meses fraccionarios, NPT de movilización sin costo, supuestos no declarados de Ecotransur) empujan todos a favor del tercerizado (Petro Star normal y Ecotransur). "
         "El único error grande en sentido contrario (precios del Lump Sum por encima de la Cot. 1255) solo afecta al Lump Sum. "
         "Verificado con el simulador: con la frecuencia y la mezcla de 2025 y las horas cotizadas por Petro Star, el ahorro de Petro Star normal pasa a −23 k USD y el de Ecotransur a −11 k USD (si tuviera los mismos tiempos): prácticamente el punto de equilibrio.")
s.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=6)
c = s.cell(rr, 2, concl); c.font = B; c.alignment = WR; s.row_dimensions[rr].height = 90
wb.save(OUT)
print("secciones", len(T) - 1, "filas sesgo", len(t["rows"]))
