"""Prueba de humo de la pila: cada herramienta se ejercita con un mini caso A vs B (datos ficticios)."""
import json, os, subprocess, sys
import numpy as np

OUT = sys.argv[1]; RECALC = sys.argv[2]
os.makedirs(OUT, exist_ok=True)
sys.path.insert(0, os.path.dirname(__file__))
from pert import muestrear, pert_media
res = {}

# 1. XlsxWriter: mini modelo con nombres definidos, fórmulas y gráfico nativo
import xlsxwriter
x = os.path.join(OUT, "mini_modelo.xlsx")
wb = xlsxwriter.Workbook(x); ws = wb.add_worksheet("Supuestos")
datos = [("Distancia_km", 120), ("Tarifa_LS_km", 900), ("Tarifa_LS_dia", 25000), ("Dias_LS", 4),
         ("Cargas", 40), ("Tarifa_carga", 3500), ("Standby_dia", 18000), ("Dias_B", 5), ("Supervision", 12000)]
for i, (n, v) in enumerate(datos):
    ws.write(i, 0, n); ws.write(i, 1, v); wb.define_name(n, f"=Supuestos!$B${i+1}")
ws.write(10, 0, "Costo_A"); ws.write_formula(10, 1, "=Distancia_km*Tarifa_LS_km+Tarifa_LS_dia*Dias_LS")
ws.write(11, 0, "Costo_B"); ws.write_formula(11, 1, "=Cargas*Tarifa_carga+Standby_dia*Dias_B+Supervision")
ws.write(12, 0, "Diferencia_B_menos_A"); ws.write_formula(12, 1, "=B12-B11")
ch = wb.add_chart({"type": "column"}); ch.add_series({"categories": "=Supuestos!$A$11:$A$12", "values": "=Supuestos!$B$11:$B$12"})
ws.insert_chart("D2", ch); wb.close()
esperado_A = 120*900 + 25000*4; esperado_B = 40*3500 + 18000*5 + 12000

# 2. Recalc con LibreOffice
r = json.loads(subprocess.run([sys.executable, RECALC, x, "60"], capture_output=True, text=True).stdout)
res["recalc_libreoffice"] = f'{r["status"]}, {r["total_formulas"]} fórmulas, {r["total_errors"]} errores'
from openpyxl import load_workbook
v = load_workbook(x, data_only=True)["Supuestos"]
assert v["B11"].value == esperado_A and v["B12"].value == esperado_B, (v["B11"].value, v["B12"].value)
res["openpyxl_lee_valores"] = f'A={v["B11"].value:,.0f}  B={v["B12"].value:,.0f}'

# 3. formulas: segundo motor
import formulas
sol = formulas.ExcelModel().loads(x).finish().calculate()
cel = {k.upper(): val for k, val in sol.items()}
b = next(val for k, val in cel.items() if k.endswith("SUPUESTOS'!B12"))
assert float(np.ravel(b.value)[0]) == esperado_B
res["formulas_segundo_motor"] = f"B={float(np.ravel(b.value)[0]):,.0f} (coincide)"

# 4. Monte Carlo PERT + SALib + AHPy + numpy-financial
rng = np.random.default_rng(7); n = 20000
dias_b = muestrear(4, 5, 9, n, rng); standby = muestrear(15000, 18000, 22000, n, rng)
costo_b = 40*3500 + standby*dias_b + 12000
p10, p50, p90 = np.percentile(costo_b, [10, 50, 90])
res["montecarlo_pert"] = f"B P10={p10:,.0f} P50={p50:,.0f} P90={p90:,.0f}; P(B<A)={np.mean(costo_b < esperado_A):.0%}; media PERT días={pert_media(4,5,9):.2f}"
from SALib.sample import morris as ms
from SALib.analyze import morris as ma
prob = {"num_vars": 3, "names": ["dias_b", "standby", "tarifa_carga"], "bounds": [[4, 9], [15000, 22000], [3000, 4200]]}
Xs = ms.sample(prob, 200, seed=1); Y = 40*Xs[:, 2] + Xs[:, 1]*Xs[:, 0] + 12000
mu = ma.analyze(prob, Xs, Y, seed=1)["mu_star"]
res["salib_ranking"] = " > ".join(n for _, n in sorted(zip(mu, prob["names"]), reverse=True))
import ahpy
c = ahpy.Compare("Criterios", {("costo", "riesgo"): 2, ("costo", "control"): 3, ("riesgo", "control"): 2}, precision=3)
res["ahpy_pesos"] = f"{c.target_weights}  CR={c.consistency_ratio}"
import numpy_financial as npf
res["numpy_financial_npv"] = f"{npf.npv(0.12, [-100, 40, 40, 40]):.2f}"

# 5. matplotlib + python-pptx con gráfico nativo → PDF → PNG
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
fig, ax = plt.subplots(figsize=(6, 3.5)); ax.hist(costo_b, bins=60, color="#4472C4")
for p in (p10, p50, p90): ax.axvline(p, color="k", ls="--", lw=1)
ax.set_title("Costo B (Monte Carlo, ficticio)"); png = os.path.join(OUT, "hist.png"); fig.tight_layout(); fig.savefig(png, dpi=150); plt.close(fig)
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = "Prueba: costo A vs B (datos ficticios)"
cd = CategoryChartData(); cd.categories = ["A Lump sum", "B Tercerizado (P50)"]; cd.add_series("Costo", (esperado_A, float(p50)))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(6), Inches(5), cd)
s.shapes.add_picture(png, Inches(6.8), Inches(1.8), width=Inches(6))
pp = os.path.join(OUT, "mini_deck.pptx"); prs.save(pp)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT, pp], capture_output=True, timeout=120)
from pdf2image import convert_from_path
imgs = convert_from_path(os.path.join(OUT, "mini_deck.pdf"), dpi=60); imgs[0].save(os.path.join(OUT, "mini_deck_p1.png"))
res["pptx_pdf_png"] = f"{len(imgs)} diapositiva(s) renderizada(s)"

# 6. markitdown: cruce de cifras
from markitdown import MarkItDown
md = MarkItDown()
t_ppt, t_xls = md.convert(pp).text_content, md.convert(x).text_content
ok = f"{esperado_A}.0" in t_ppt and f"{esperado_A}.0" in t_xls
res["markitdown_cruce"] = f"costo A {esperado_A:,} aparece en deck y Excel: {'OK' if ok else 'FALLA'}"
for k, v in res.items(): print(f"{k:24s} {v}")
