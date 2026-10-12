"""Agrega al simulador las hojas: Método, Reconciliación, Medido y supuesto, Sensibilidad, Pendientes."""
import sys, json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL

PATH, SENS = sys.argv[1], sys.argv[2]
S = json.load(open(SENS))
wb = load_workbook(PATH)
F = "Arial"
B = Font(name=F, size=10, bold=True); N = Font(name=F, size=10); NOTE = Font(name=F, size=9, italic=True, color="595959")
TITLE = Font(name=F, size=14, bold=True, color="1F3864"); H1 = Font(name=F, size=11, bold=True, color="FFFFFF")
GRN = Font(name=F, size=10, color="008000"); BLUE = Font(name=F, size=10, color="0000FF")
F_H1 = PatternFill("solid", fgColor="1F3864"); F_H2 = PatternFill("solid", fgColor="D9E1F2"); F_Y = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '$#,##0;($#,##0);"-"'; USD2 = '$#,##0.00;($#,##0.00);"-"'
WR = Alignment(wrap_text=True, vertical="top")


def sheet(name, widths, title, sub=None):
    s = wb.create_sheet(name); s.sheet_view.showGridLines = False
    for i, w in enumerate(widths): s.column_dimensions[CL(i + 1)].width = w
    s["B1"] = title; s["B1"].font = TITLE
    if sub: s["B2"] = sub; s["B2"].font = NOTE
    return s


def table(s, r0, headers, rows, fmts=None, height=None):
    for i, h in enumerate(headers):
        c = s.cell(r0, 2 + i, h); c.font = B; c.fill = F_H2; c.border = BOX; c.alignment = WR
    for j, row in enumerate(rows):
        for i, v in enumerate(row):
            c = s.cell(r0 + 1 + j, 2 + i, v); c.border = BOX; c.alignment = WR
            c.font = GRN if isinstance(v, str) and v.startswith("=Simulador") else N
            if fmts and fmts[i]: c.number_format = fmts[i]
        if height: s.row_dimensions[r0 + 1 + j].height = height
    return r0 + 1 + len(rows)


def h1(s, r, text, last):
    s.merge_cells(f"B{r}:{last}{r}"); c = s[f"B{r}"]; c.value = text; c.font = H1; c.fill = F_H1


# ================= MÉTODO =================
m = sheet("Método", [2, 34, 22, 22, 30, 60], "Método: cómo está armado el costo de cada modelo y qué se cancela",
          "Se lee antes del simulador. Los umbrales son el estudio; el promedio no.")
h1(m, 4, "1. CLASIFICACIÓN DE CADA RENGLÓN DE COSTO", "F")
rows = [
    ("Tarifa por movilización (Work Order)", "Work Order", "Variable por movilización; escalón por rango", "Límites de rango 0,2 / 5 / 10 km", "No es monótona: 5–10 km (95.374) cuesta menos que 0,1–5 km (98.663)."),
    ("Recargo de señalizadoras faltantes", "Work Order", "Variable; escalón 1 o 2 turnos", "12 h = 1 turno; más de 12 h = 2 turnos", "4 personas por turno × 375.000 COP (Cot. 1255 p. 5)."),
    ("Tarifa del transportador", "Tercerizado", "Variable; escalón por rango", "0,2 / 5 / 10 / 20 km", "Petro Star no cotiza más de 20 km: acantilado de cobertura."),
    ("Descuento desde la 6ª del mes", "Tercerizado", "Acantilado por número de movilización del mes", "6ª movilización del mes", "Petro Star: 50 % de la tarifa. Ecotransur: solo equipos. Lump Sum: sin descuento (evaluación principal)."),
    ("Cargo fijo mensual (nómina + disponibilidad)", "Tercerizado", "Escalón decreciente; fijo mensual", "Petro Star 0–5 movs; Ecotransur 0–4 movs", "Petro Star todavía cobra con 5 movs (p. 4); Ecotransur no cobra desde la 5ª."),
    ("Taladro durante la movilización", "Tercerizado", "Variable por hora", "Horas por rango (12/18/24/28)", "34.944 USD/día. Bajo el WO no se paga (premisa SierraCol)."),
    ("Truck Manager", "Tercerizado", "Fijo mensual en la realidad; variable por mov. en el Excel", "60 M COP/mes ÷ 6 movs", "Por debajo de 6 movs/mes el prorrateo del Excel lo subestima; por encima lo sobreestima."),
    ("Horas excedidas sobre el tiempo declarado", "Tercerizado (Petro Star)", "Acantilado", "Tiempo total 12/24/36/36 h", "Cot. 1255 p. 5 y cláusula 7: equipos y nómina por hora si la demora no es atribuible al transportador."),
    ("Bandas de horas por ítem (Opción A/B)", "Tercerizado (Ecotransur)", "Acantilado", "8–12 / ≤14 / 14–18 / 18–24 / ≤22 h", "El formato del Ítem 1 no viene adjunto: el Excel usa ítems 'más de X horas'."),
    ("Personal, turnos o equipos fuera del mínimo", "Ambos", "Variable", "Lo que exceda el equipo mínimo", "Premisa: se cancela entre modelos (adicionales 0 %). Ver Pendientes."),
]
r = table(m, 5, ["Renglón", "Modelo", "Tipo", "Umbral", "Nota"], rows, height=42)
h1(m, r + 1, "2. ¿QUÉ VARIABLE AFECTA A LOS DOS MODELOS Y QUÉ SE CANCELA?", "F")
rows = [
    ("Número de pozos de la campaña", "Ambos", "No", "El WO crece lineal; el tercerizado tiene umbrales mensuales. Entra completo al diferencial."),
    ("Frecuencia mensual de movilizaciones", "Solo tercerizado", "No", "El WO cobra por movilización, no por mes. Es la variable que define la frontera."),
    ("Rango de distancia", "Ambos, distinto", "No", "Las tarifas no se mueven igual: pasar de 0,1–5 a 5–10 km abarata el WO y encarece el tercerizado."),
    ("Adicionales / imprevistos", "Ambos", "Sí (premisa)", "Se pone 0 % en ambos. Si no se cancelan, el residuo cae sobre el tercerizado (ver Sensibilidad)."),
    ("Taladro parado por NPT de perforación", "Ambos", "Sí", "La tarifa del taladro durante la perforación la rige su contrato en los dos modelos."),
    ("Demoras dentro de la movilización", "Solo tercerizado", "No", "El WO es suma cerrada. En el tercerizado se paga el taladro por hora y, con Petro Star, equipos y nómina (cláusula 7)."),
    ("TRM", "Solo tercerizado", "No", "Las ofertas están en COP; el WO en USD. TRM alta favorece al tercerizado."),
    ("Tarifa del taladro", "Solo tercerizado", "No", "Bajo el WO no se pagan horas de taladro."),
    ("Señalizadoras", "Solo WO (recargo)", "No", "Las ofertas ya incluyen 20; el WO tiene 12."),
    ("Truck Manager", "Solo tercerizado", "No", "Premisa SierraCol."),
]
r = table(m, r + 2, ["Variable", "¿A quién afecta?", "¿Se cancela?", "Por qué"], rows, height=30)
h1(m, r + 1, "3. NPT: POR QUÉ CANAL AFECTA A CADA MODELO", "F")
txt = [
    "NPT de perforación (pozo): el taladro está parado o perforando más días. Su tarifa diaria se paga igual en los dos modelos, así que ese costo se cancela. Lo que NO se cancela es el efecto sobre la frecuencia: menos movilizaciones por mes. Al WO no le pasa nada (paga por movilización). Al tercerizado le cambia la factura: si el mes cae por debajo de 5 (Ecotransur) o de 6 (Petro Star), aparece el cargo fijo y se pierde el descuento. Canal: frecuencia, no horas.",
    "NPT de movilización (demora en el traslado): en el WO lo absorbe el contratista del taladro (suma global cerrada). En el tercerizado carga horas de taladro (1.456 USD/h) y, con Petro Star, si la demora no es atribuible al transportador, también equipos y nómina (Cot. 1255 p. 5, cláusula 7; a tarifas horarias de la p. 5 la flota mínima más 20 señalizadores suma unos 1.357 USD/h). Ecotransur no trae esa cláusula, pero factura aparte turnos extra. Canal: horas, residuo completo sobre el tercerizado.",
    "Lo único medido en los archivos es el sobretiempo de movilización de los logs (hojas Real 2024 / Real 2025): misma localización excede las 12 h del Excel en 10 de 23 movs (2024) y 10 de 11 (2025); 0,1–5 km excede las 18 h en 11 de 28 y 10 de 17; 5–10 km y más de 10 km en 2024 quedaron por debajo de lo supuesto. No hay NPT de perforación ni plan de pozos: se deja editable en 0 % (Simulador C23) y se muestra su sensibilidad.",
]
for i, t in enumerate(txt):
    rr = r + 2 + i; m.merge_cells(f"B{rr}:F{rr}"); c = m[f"B{rr}"]; c.value = t; c.font = N; c.alignment = WR; m.row_dimensions[rr].height = 66
r = r + 2 + len(txt)
h1(m, r + 1, "4. VERIFICACIONES OBLIGATORIAS", "F")
rows = [
    ("Recalcular fórmulas heredadas", "Hecho", "Las filas amarillas 'Costo Suma Global' (31, 57, 83, 109) se reproducen al centavo (hoja Reconciliación). El recargo de señalizadoras da 405,41 / 810,81 USD. El total WO de la campaña es 3.130.187 USD (= M18); la nota J55 cita 3.107.890."),
    ("Monotonía de tarifas con la distancia", "Hallazgo", "Petro Star, Lump Sum y Ecotransur crecen con la distancia. El WO no: 5–10 km < 0,1–5 km. Si los 20 pozos fueran de 5–10 km, el ahorro de Petro Star cae unos 324.000 USD y el de Ecotransur unos 430.000 USD."),
    ("Cobertura de más de 20 km", "Brecha", "Petro Star cotiza 10–20 km. El log 2024 tiene 1 movimiento de 22,1 km (pozo 5681); 2025 ninguno (máx. 14,6 km). Se valoró con la tarifa de 10–20 km y queda marcado."),
    ("Horas de Petro Star vs Excel", "Diferencia", "Cotización: 12 / 24 / 36 / 36 h. Excel: 12 / 18 / 24 / 28 h. Por movilización: 0 / 8.736 / 17.472 / 11.648 USD de taladro. En la campaña base (20 pozos en 0,1–5 km): 174.720 USD. Editable en Simulador fila 39."),
    ("IVA del Lump Sum", "Fuera de alcance", "Por instrucción de SierraCol se trabaja sin IVA. Si el IVA del 50 % de alquiler no fuera descontable, el Lump Sum costaría unos 225.000 USD más en la campaña base."),
    ("Celdas sueltas F155 y F156", "No usadas", "Valores 8 y 15 sin rótulo, dentro del bloque 5.6. No entran en ningún cálculo de este simulador."),
]
table(m, r + 2, ["Verificación", "Resultado", "Detalle"], rows, height=54)

# ================= RECONCILIACIÓN =================
rc = sheet("Reconciliación", [2, 46, 16, 16, 12, 52, 20], "Reconciliación: de dónde sale cada parámetro",
           "Valor en simulador = enlace vivo (verde). Valor en fuente = copiado del archivo original. Control = OK si coinciden; el control solo es válido con los valores por defecto del simulador.")
rows = []
def add(lab, ref, val, src, cls, fmt=None):
    rows.append((lab, f"=Simulador!{ref}", val, f'=IF(ABS(C{5+len(rows)}-D{5+len(rows)})<0.005,"OK","DIFERENCIA")', src, cls))
add("TRM (COP/USD)", "C7", 3700, "Tarifas.xlsx · E3", "Medido")
add("Taladro (USD/día)", "C8", 34944, "Tarifas.xlsx · E5", "Medido")
add("Taladro (USD/h)", "C9", 1456, "Tarifas.xlsx · E6", "Medido")
add("Adicionales base (%)", "C11", 0, "Tarifas.xlsx · E7", "Premisa")
add("Adicionales Lump Sum (%)", "C12", 0, "Tarifas.xlsx · E8", "Premisa")
add("Truck Manager por movilización (USD)", "C54", 2702.7027027027, "Tarifas.xlsx · C29", "Medido")
for g, (col, rr) in enumerate(zip("CDEF", [26, 52, 78, 104])):
    add(f"Horas rango {g+1}", f"{col}39", [12, 18, 24, 28][g], f"Tarifas.xlsx · C{rr}", "Dato Excel")
for g, col in enumerate("CDEF"):
    add(f"WO base rango {g+1} (USD)", f"{col}43", [75642, 98663, 95374, 110174][g], f"Tarifas.xlsx · K{4+g} (dato SierraCol)", "Premisa")
for g, (col, rr) in enumerate(zip("CDEF", [25, 51, 77, 103])):
    add(f"Petro Star normal rango {g+1} (COP)", f"{col}44", [140385441, 220853084, 238462998, 246462998][g], f"Tarifas.xlsx · E{rr}", "Medido")
    add(f"Lump Sum rango {g+1} (COP)", f"{col}45", [178363885, 304284690, 373649874, 414800928][g], f"Tarifas.xlsx · F{rr} (≠ Cot. 1255 p.1–2)", "Medido")
    add(f"Ecotransur rango {g+1} (COP)", f"{col}46", [140233002, 207703002, 249003002, 253060002][g], f"Tarifas.xlsx · G{rr}", "Medido")
add("Señalizador turno 12 h (COP)", "C21", 375000, "Cot. 1255 p. 5", "Medido")
add("Recargo señalizadoras misma loc. (USD)", "C50", 405.405405405405, "Tarifas.xlsx · K11 (verificado)", "Cálculo")
add("Recargo señalizadoras otros rangos (USD)", "D50", 810.810810810811, "Tarifas.xlsx · L11 (verificado)", "Cálculo")
yel = {"PS": ([58116.71378378378, 88600.7254054054, 102096.16162162162, 110082.3237837838], 62, "E"),
       "LS": ([68381.15810810811, 111149.80810810812, 138633.15513513514, 155579.06162162163], 64, "F"),
       "EC": ([58075.51405405405, 85046.6491891892, 104944.81135135135, 111865.29783783785], 66, "G")}
for o, (vals, row, xc) in yel.items():
    for g, col in enumerate("CDEF"):
        add(f"Costo Suma Global {o} rango {g+1} (USD)", f"{col}{row}", vals[g], f"Tarifas.xlsx · {xc}{[31,57,83,109][g]} (fila amarilla)", "Control")
add("WO de la campaña 15+20 (USD)", "C83", 3130187.2972973, "Tarifas.xlsx · M18 (la nota J55 dice 3.107.890)", "Control")
for i, v in enumerate([430000000, 164780000, 327959522, 166920000]):
    pass
r = table(rc, 4, ["Parámetro", "Valor en simulador", "Valor en fuente", "Control", "Fuente (archivo · celda o página)", "Clase"], rows, [None, USD2, USD2, None, None, None])
rc.cell(r + 1, 2, "Cargos fijos mensuales: Simulador C72:D78 transcritos de Cot. 1255 p. 4 (nómina + disponibilidad) y Ecotransur p. 2–3 (nómina + stand-by). Equipos Ecotransur 6ª+: p. 3–4. Perfiles de frecuencia: logs C11:C21 (2024) y C11:C19 (2025).").font = NOTE

# ================= MEDIDO Y SUPUESTO =================
ms = sheet("Medido y supuesto", [2, 50, 30, 60], "Lo medido y lo supuesto, separados",
           "Ningún supuesto tiene la autoridad de un dato medido. Los supuestos están en amarillo en el simulador.")
med = [("Tarifas de las tres ofertas por rango", "Tarifas.xlsx E/F/G", "Precios tal cual SCM los pasó (F no coincide con la Cot. 1255)."),
       ("TRM y tarifa del taladro", "Tarifas.xlsx E3, E5", ""),
       ("Cargo fijo mensual por número de movilizaciones", "Cot. 1255 p. 4 · Ecotransur p. 2–3", ""),
       ("Descuento desde la 6ª (Petro Star 50 %; Ecotransur solo equipos)", "Cot. 1255 p. 2 · Ecotransur p. 3–4", ""),
       ("Tarifa del señalizador (375.000 COP/turno)", "Cot. 1255 p. 5", ""),
       ("Frecuencia mensual 2024 y 2025", "Logs, tabla por mes", "Dos observaciones independientes; no se promedian."),
       ("Mezcla por rango 2024 (23/29/14/4) y 2025 (11/17/10/7)", "Logs, resumen por rango", ""),
       ("Horas reales por movilización", "Logs, columna Tiempo mov hr", "Pozo 5136 sin horas."),
       ("Horas totales declaradas por Petro Star (12/24/36/36)", "Cot. 1255 p. 1–2", "Referencia; el simulador usa las del Excel."),
       ("Cláusula de demoras no atribuibles (Petro Star)", "Cot. 1255 p. 5, condición 7", "Ecotransur no la trae.")]
sup = [("Work Order base y 12 señalizadoras", "Premisa SierraCol", "No está en los archivos."),
       ("Adicionales 0 % en ambos modelos", "Premisa SierraCol", "Sensibilidad: +5 % solo en tercerizado."),
       ("Truck Manager por movilización (60 M ÷ 6)", "Premisa SierraCol", "En la realidad es mensual."),
       ("Campaña 15 misma loc. + 20 otra loc.", "Premisa SierraCol", "No sale de los logs."),
       ("Rango de los 20 pozos: 0,1–5 km", "Supuesto", "Los logs muestran 5–10 y más de 10 km en 38–50 % de los movimientos fuera de la locación."),
       ("Lump Sum sin descuento (principal)", "Premisa SierraCol", "Con descuento: solo referencia."),
       ("Ecotransur 6ª+ en misma loc. = equipos 0–5 km", "Supuesto", "El PDF no trae valor."),
       ("Orden de los pozos dentro del mes: intercalado", "Supuesto", "Sensibilidad: misma loc. primero / otra primero."),
       ("Horas pagadas al taladro = horas del Excel", "Dato del Excel", "Los logs muestran más horas en misma loc. y 0,1–5 km (2025)."),
       ("NPT de perforación 0 %", "Supuesto (sin dato)", "Barrido 0–20 % en Sensibilidad."),
       ("Meses sin registro = campaña parada (no se cobran)", "Premisa SierraCol", "Si el contrato sigue vigente: ver Pendientes."),
       ("Más de 20 km se paga con la tarifa de 10–20 km", "Supuesto", "1 movimiento en 2024."),
       ("IVA fuera del análisis", "Premisa SierraCol", "")]
r = table(ms, 4, ["MEDIDO", "Fuente", "Nota"], med, height=28)
table(ms, r + 2, ["SUPUESTO O PREMISA", "Origen", "Nota"], sup, height=28)

# ================= SENSIBILIDAD =================
se = sheet("Sensibilidad", [2, 46, 18, 22, 26, 16, 26, 16, 16], "Sensibilidad: qué variable mueve el ahorro de la campaña base",
           "Ahorro = costo WO − costo tercerizado (USD), campaña 15 misma loc. + 20 en 0,1–5 km. 'Mejora' = extremo que aumenta el ahorro del tercerizado; 'Empeora' = el que lo reduce. Ordenado por efecto. Calculado con el motor de referencia (tools/modelo.py), idéntico al simulador; no se recalcula al editar.")
h1(se, 4, "ESCENARIOS COMBINADOS (campaña 15 + 20; ahorro frente al WO, USD)", "I")
SC = [tuple(x) for x in S["escenarios"]]
r = table(se, 5, ["Escenario", "Petro Star normal", "Lump Sum (sin desc.)", "Ecotransur"], SC, [None, USD, USD, USD], height=28)
se.cell(r, 2, "Horas cotizadas se aplican igual a Ecotransur (supuesto: su carta no declara tiempos). Calculado con tools/modelo.py.").font = NOTE
r += 2
names = {"PS": "Petro Star tarifa normal", "LS": "Petro Star Lump Sum (sin descuento)", "EC": "Ecotransur"}
for prof in ("2024", "2025"):
    for o in ("PS", "LS", "EC"):
        rows_o = sorted([x for x in S[prof]["rows"] if x["oferta"] == o and x["efecto_usd"] > 0.5], key=lambda x: -x["efecto_usd"])
        h1(se, r, f"{names[o]} · frecuencia perfil {prof} · ahorro base {S[prof]['base'][o]:,.0f} USD".replace(",", "."), "I")
        data = [(x["variable"], x["clase"], x["base"], x["favorable_txt"], x["favorable_usd"], x["desfavorable_txt"], x["desfavorable_usd"], x["efecto_usd"]) for x in rows_o]
        r = table(se, r + 1, ["Variable", "Clase", "Valor base", "Extremo que mejora", "Ahorro (USD)", "Extremo que empeora", "Ahorro (USD)", "Efecto (USD)"], data,
                  [None, None, None, None, USD, None, USD, USD], height=28) + 1

# ================= PENDIENTES =================
pe = sheet("Pendientes", [2, 44, 26, 30, 16, 16, 16, 16, 30], "Pendientes por confirmar: valor de trabajo y lo que está en juego",
           "En juego = cambio del ahorro del tercerizado (USD) si el valor de trabajo se reemplaza por la alternativa. Negativo = el tercerizado pierde ventaja. Campaña base 15 + 20; perfil de frecuencia 2024 / 2025. Valores fijos calculados con tools/sensibilidad.py (mismo motor que el simulador); no se recalculan al editar el simulador.")
P = {p["perfil"]: p for p in S["pendientes"]}
def st(k, o): return (P["2024"][k][o], P["2025"][k][o])
rows = []
def addp(item, trabajo, alt, k, o, quien, ofname):
    a, b = st(k, o); rows.append((item, trabajo, alt, ofname, a, b, quien))
addp("Rango de los 20 pozos 'Dif. Loc'", "0,1–5 km", "5–10 km", "rango_5_10", "PS", "SierraCol (plan de pozos)", "Petro Star normal")
addp("Rango de los 20 pozos 'Dif. Loc'", "0,1–5 km", "5–10 km", "rango_5_10", "EC", "SierraCol (plan de pozos)", "Ecotransur")
addp("Rango de los 20 pozos 'Dif. Loc'", "0,1–5 km", "más de 10 km", "rango_10", "PS", "SierraCol (plan de pozos)", "Petro Star normal")
addp("Rango de los 20 pozos 'Dif. Loc'", "0,1–5 km", "más de 10 km", "rango_10", "EC", "SierraCol (plan de pozos)", "Ecotransur")
addp("Horas que se pagan al taladro", "Excel 12/18/24/28", "Petro Star 12/24/36/36", "horas_ps", "PS", "SierraCol / Petro Star", "Todas (igual)")
addp("Petro Star normal: ¿cobra aparte los equipos adicionales requeridos?", "No", "Sí (Cot. 1255, azul)", "ps_adic", "PS", "Petro Star", "Petro Star normal")
addp("Lump Sum: valores del Excel vs Cot. 1255", "Excel (F)", "Cot. 1255 p. 1–2", "LS_pdf", "LS", "SCM", "Lump Sum")
addp("Lump Sum con descuento desde la 6ª", "Sin descuento (principal)", "50 % desde la 6ª", "LS_desc", "LS", "Petro Star", "Lump Sum")
addp("Petro Star: ¿cobra cargo fijo con 5 movs?", "Sí (tabla p. 4)", "No (mínimo 5 cumplido)", "fijo5", "PS", "Petro Star", "Petro Star (ambas)")
addp("Ecotransur 6ª+ en misma localización", "Equipos 0–5 km", "Tarifa plena", "ec_misma", "EC", "Ecotransur", "Ecotransur")
addp("Truck Manager", "Por movilización", "60 M COP fijo por mes", "tm_mensual", "PS", "SierraCol", "Todas (igual)")
addp("Adicionales: ¿se cancelan?", "0 % ambos", "+5 % solo tercerizado", "adic5", "PS", "SierraCol", "Petro Star normal")
addp("Adicionales: ¿se cancelan?", "0 % ambos", "+5 % solo tercerizado", "adic5", "EC", "SierraCol", "Ecotransur")
addp("Recargo de señalizadoras al WO", "Se carga", "No se carga", "senal_no", "PS", "SierraCol (contrato WO)", "Todas (igual)")
rows.append(("Mes con campaña parada y contrato vigente (por mes)", "No se cobra", "Se cobra el cargo de 0 movs", "Petro Star (ambas)", -S["mes_parado_usd"]["PS"], -S["mes_parado_usd"]["PS"], "Contrato (cláusula de suspensión)"))
rows.append(("Mes con campaña parada y contrato vigente (por mes)", "No se cobra", "Se cobra el cargo de 0 movs", "Ecotransur", -S["mes_parado_usd"]["EC"], -S["mes_parado_usd"]["EC"], "Contrato (cláusula de suspensión)"))
rows.append(("Demora no atribuible al transportador (por hora de demora, igual en ambos perfiles)", "Sin demora", "1 hora de demora", "Petro Star", -2813.2, -2813.2, "Cot. 1255 condición 7: taladro 1.456 + flota mínima y 20 señalizadores ~1.357 USD/h (tarifas p. 5)"))
rows.append(("Perfil 2025 empieza en febrero con 1 movilización (¿mes parcial?)", "Cuenta como mes de campaña", "No cuenta", "Todas", None, None, "SierraCol (fecha real de inicio de la campaña 2025)"))
rows.append(("Ecotransur 6ª+: ¿cobra además el personal operativo usado (señaleras, supervisión)?", "No (solo equipos del Ítem 2)", "Sí", "Ecotransur", None, None, "Ecotransur (carta p. 4: 'equipos y el personal operativo efectivamente utilizado')"))
rows.append(("Riesgos no cuantificados del tercerizado: coordinación taladro–transportador, responsabilidad HSE y por daños (pólizas), disponibilidad del contratista, plazo y salida del contrato, pago en COP a 30 días y cargo fijo mensual", "Sin valorar", "—", "Todas", None, None, "Contratos / HSE / Tesorería"))
rows.append(("Movilizaciones de más de 20 km", "Tarifa 10–20 km", "Sin tarifa", "Petro Star", None, None, "Petro Star (cotizar)"))
rows.append(("Formato Ítem 1 y 2 de Ecotransur (bandas de horas A/B)", "Excel ítems 1.3/1.9/1.15/1.21", "Ítem según horas reales", "Ecotransur", None, None, "Ecotransur (enviar formato)"))
rows.append(("NPT de perforación", "0 %", "20 % (barrido)", "Petro Star normal",
             [x for x in S["2024"]["rows"] if x["oferta"] == "PS" and x["variable"].startswith("NPT")][0]["desfavorable_usd"] - S["2024"]["base"]["PS"],
             [x for x in S["2025"]["rows"] if x["oferta"] == "PS" and x["variable"].startswith("NPT")][0]["desfavorable_usd"] - S["2025"]["base"]["PS"], "SierraCol (plan de pozos / NPT)"))
table(pe, 4, ["Pendiente", "Valor de trabajo", "Alternativa", "Oferta afectada", "En juego perfil 2024 (USD)", "En juego perfil 2025 (USD)", "Quién confirma"], rows,
      [None, None, None, None, USD, USD, None], height=30)

wb.move_sheet("Método", offset=-(wb.sheetnames.index("Método") - 1))
wb.save(PATH)
print(wb.sheetnames)
